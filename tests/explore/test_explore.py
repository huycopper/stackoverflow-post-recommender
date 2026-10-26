from __future__ import annotations

from datetime import UTC, datetime

import pyarrow as pa
import pytest
from openpyxl import load_workbook

from so_recommender.collect.raw_store import RawStore
from so_recommender.explore.excel import build_workbook
from so_recommender.explore.flatten import flatten
from so_recommender.explore.stats import bin_counts, compute, duplicate_groups, load

from ..collect.conftest import FakeClock
from ..collect.test_collector import build_api, make_collector, question


@pytest.fixture
def flat(tmp_path):
    """Collect the fake API into tmp/raw, add awkward extra raw records, flatten to tmp/flat."""
    collector = make_collector(build_api(), FakeClock(), tmp_path)
    assert collector.run() is True
    raw = RawStore(tmp_path / "raw", clock=lambda: datetime(2030, 1, 1, tzinfo=UTC))  # later
    # the same question fetched again (e.g. after a resume) must appear once
    raw.write(
        "questions_linked", "u", {}, {"items": [question(10, comments=2, dup_of=(1, 50, 999))]}
    )
    # background question 3 fetched again after it was closed as a duplicate: latest copy wins
    raw.write("questions_by_id", "u", {}, {"items": [question(3, dup_of=(1,))]})
    # a title that looks like a spreadsheet formula
    raw.write(
        "questions",
        "u",
        {},
        {"items": [question(77, tags=("python", "excel")) | {"title": "=1+1"}]},
    )
    counts = flatten(tmp_path / "raw", tmp_path / "flat", collector.state, chunk_rows=2)
    return tmp_path / "flat", counts


def test_flatten_keeps_the_latest_copy_of_each_record(flat):
    flat_dir, counts = flat
    assert counts == {"questions": 7, "answers": 1, "comments": 4, "duplicate_pairs": 5}
    t = load(flat_dir)
    q = {r["question_id"]: r for r in t["questions"].to_pylist()}
    assert q[10]["source"] == "questions_linked" and q[10]["original_ids"] == [1, 50, 999]
    assert q[3]["closed_reason"] == "Duplicate" and q[3]["source"] == "questions_by_id"
    assert q[1]["is_seed"] and q[3]["is_background"] and not q[10]["is_seed"]
    assert pa.types.is_timestamp(t["questions"].schema.field("creation_date").type)
    comments = {r["comment_id"]: r for r in t["comments"].to_pylist()}
    assert comments[600]["post_type"] == "answer" and comments[600]["question_id"] == 1
    pairs = set(
        zip(
            t["duplicate_pairs"]["duplicate_id"].to_pylist(),
            t["duplicate_pairs"]["original_id"].to_pylist(),
            strict=True,
        )
    )
    assert pairs == {(2, 1), (3, 1), (10, 1), (10, 50), (10, 999)}


def test_duplicate_groups_and_bins():
    groups = duplicate_groups([(1, 2), (3, 2), (4, 5)])
    assert sorted(sorted(m) for m in groups.values()) == [[1, 2, 3], [4, 5]]
    assert bin_counts([0, 1, 1, 7], [("0", 0, 0), ("1", 1, 1), ("> 1", 2, None)]) == [
        ["0", 1],
        ["1", 2],
        ["> 1", 1],
    ]


def test_stats_on_flat_tables(flat):
    stats = compute(load(flat[0]))
    assert stats["overview"]["Câu hỏi"] == 7
    assert stats["overview"]["Nhóm trùng (thành phần liên thông)"] == 1
    assert stats["roles"]["bị đóng vì trùng (query)"] == 3
    per_q = dict(stats["comments_per_question"]["all"])
    assert per_q["0"] == 5 and per_q["1"] == 1 and per_q["2"] == 1
    assert stats["leaks"]["any_comment"] == 1
    assert stats["leaks"]["duplicates_with_leaky_comment"] == 1
    assert stats["biggest_groups"][0][:3] == [6, 1, 3]  # question 1 is the hub of 2, 3 and 10


def test_workbook_sheets_formulas_and_text_safety(flat, tmp_path):
    tables = load(flat[0])
    out = build_workbook(tables, compute(tables), tmp_path / "x.xlsx", n_questions=50)
    wb = load_workbook(out)
    assert wb.sheetnames == [
        "Tổng quan",
        "Thống kê",
        "Cặp trùng (mẫu)",
        "Comment lộ đáp án (mẫu)",
        "Comment ngẫu nhiên (mẫu)",
        "Câu hỏi (mẫu)",
    ]
    formulas = [c.value for ws in wb for row in ws.iter_rows() for c in row if c.data_type == "f"]
    assert formulas and all(f.startswith("=IFERROR(") for f in formulas)
    titles = [c.value for c in wb["Câu hỏi (mẫu)"]["C"][1:]]
    assert "=1+1" in titles  # stored as text, not as a formula
    assert wb["Cặp trùng (mẫu)"].max_row == 1 + 4  # pairs whose original is in the data
